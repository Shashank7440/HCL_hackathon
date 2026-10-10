"""Parse files to pages. pdf: one Page per page (pymupdf, guarded); md/txt/docx/html -> page 1."""
from __future__ import annotations
from dataclasses import dataclass

@dataclass
class Page:
    page: int
    text: str

def parse_file(path: str, filename: str) -> list[Page]:
    lower = (filename or path).lower()
    if lower.endswith(".pdf"):
        try:
            import fitz                                   # pymupdf
            doc = fitz.open(path)
            return [Page(i + 1, p.get_text()) for i, p in enumerate(doc)]
        except ImportError:
            import pypdf
            reader = pypdf.PdfReader(path)
            return [Page(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if lower.endswith(".docx"):
        import docx as pydocx
        d = pydocx.Document(path)
        parts = []
        for p in d.paragraphs: parts.append(p.text)
        for t in d.tables:
            for row in t.rows: parts.append(" | ".join(c.text for c in row.cells))
        return [Page(1, "\n".join(parts))]
    if lower.endswith((".html", ".htm")):
        from html.parser import HTMLParser
        class _T(HTMLParser):
            def __init__(self): super().__init__(); self.buf = []
            def handle_data(self, d): self.buf.append(d)
        t = _T(); t.feed(open(path, encoding="utf-8", errors="ignore").read())
        return [Page(1, " ".join(t.buf))]
    return [Page(1, open(path, encoding="utf-8", errors="ignore").read())]     # txt, md
